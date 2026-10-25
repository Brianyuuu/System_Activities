#!/usr/bin/env python3
"""Industry matrices from validated final assignments, explicit coverage, and World memberships."""
import argparse
import collections
import csv
import hashlib
import json
import os
import re
import unicodedata
import urllib.request
from pathlib import Path
import numpy as np

CONTRACT = {'model': 'gemini-embedding-2', 'dimensions': 3072,
            'prefix': 'task: sentence similarity | query: '}
COLUMNS = ['firm', 'year', 'parents count', 'parents enter', 'parents exit',
           'parents new to world', 'closest firm', 'distance to closest firm', 'distance to avg']
COMPLETE = {'assigned', 'rejected_by_gemini', 'no_embedding_candidate', 'rejected_by_glm'}

def sha(p):
    return hashlib.sha256(p.read_bytes()).hexdigest()

def read(p):
    with p.open(encoding='utf-8-sig', newline='') as f:
        return list(csv.DictReader(f))

def dump(p, value):
    tmp = p.with_suffix(p.suffix + '.tmp')
    tmp.write_text(json.dumps(value, indent=2, ensure_ascii=False, allow_nan=False) + '\n')
    tmp.replace(p)

def normalize(label):
    return ' '.join(unicodedata.normalize('NFKC', label).casefold().split())

def prepare(run, coverage, world, world_metadata, output):
    if output.exists():
        raise FileExistsError('Choose a new matrix directory')
    final = run / 'outputs/pipeline/final'
    v = json.loads((final / 'validation.json').read_text())
    taxonomy = run / 'inputs/final_taxonomy_enrich.csv'
    original = run / 'inputs/source_original.csv'
    assignments = final / 'all_assigned.csv'
    if not v.get('complete') or v.get('unresolved') != 0:
        raise ValueError('Final assignment validation must be complete with zero unresolved rows')
    for p, digest in [(taxonomy, v['taxonomy_sha256']), (original, v['source_original_sha256']),
                      (assignments, v['output_hashes']['all_assigned.csv'])]:
        if sha(p) != digest:
            raise ValueError('Assignment input hash mismatch: ' + p.name)
    tax_rows = read(taxonomy)
    # Assignment adapter trims labels before creating its frozen parent snapshot.
    parents = {r['parent_id']: r['parent_label'].strip() for r in tax_rows}
    if len(parents) != len(tax_rows):
        raise ValueError('Duplicate taxonomy IDs')
    panel = {}
    for row in read(coverage):
        key = row['firm'], int(row['year'])
        if key in panel or row['complete'] not in {'0', '1'}:
            raise ValueError('Duplicate coverage row or invalid complete flag')
        panel[key] = row
    if not panel:
        raise ValueError('Empty coverage manifest')
    source, assigned = read(original), read(assignments)
    if len(source) != len(assigned) or len(source) != v['source_rows']:
        raise ValueError('Source row conservation failed')
    seen, counts, hits = set(), collections.Counter(), collections.defaultdict(set)
    for row in assigned:
        index = int(row['source_original_row']) - 2
        if index in seen or not 0 <= index < len(source):
            raise ValueError('Invalid source row mapping')
        seen.add(index)
        if any(row.get(k) != value for k, value in source[index].items()):
            raise ValueError('Original source fields changed')
        key = row['ticker'], int(row['year'])
        if key not in panel or row['assignment_status'] not in COMPLETE:
            raise ValueError('Out-of-scope or unfinished assignment')
        counts[key] += 1
        ids, labels = json.loads(row['parent_ids_json']), json.loads(row['parent_labels_json'])
        if len(ids) != len(labels) or len(ids) != int(row['parent_count']) or len(set(ids)) != len(ids):
            raise ValueError('Parent count or uniqueness mismatch')
        if bool(ids) != (row['assignment_status'] == 'assigned'):
            raise ValueError('Assignment status conflicts with confirmed parents')
        for pid, label in zip(ids, labels):
            if parents.get(pid) != label:
                raise ValueError('Parent ID/label mismatch')
            if panel[key]['complete'] == '1':
                hits[key].add(label)
    for key, row in panel.items():
        if int(row['expected_source_rows']) < 0 or counts[key] != int(row['expected_source_rows']):
            raise ValueError('Expected source count mismatch: ' + str(key))
    metadata = json.loads(world_metadata.read_text())
    if metadata.get('taxonomy_sha256') != sha(taxonomy) or metadata.get('world_sha256') != sha(world):
        raise ValueError('World provenance hash mismatch')
    if metadata.get('basis') not in {'confirmed_assignments', 'current_catalog_memberships'} or not metadata.get('scope_description'):
        raise ValueError('World metadata needs a declared validity basis and scope')
    earliest = {}
    for row in read(world):
        key = row['firm'], int(row['year'])
        if not row['source_id'] or parents.get(row['parent_id']) != row['parent_label']:
            raise ValueError('Untraceable World membership or parent mismatch')
        if key in panel and panel[key]['complete'] == '0':
            continue
        label = row['parent_label']
        earliest[label] = min(earliest.get(label, key[1]), key[1])
    for (firm, year), labels in hits.items():
        for label in labels:
            earliest[label] = min(earliest.get(label, year), year)
    records = []
    for key in sorted(panel):
        firm, year = key; complete = panel[key]['complete'] == '1'; labels = hits[key]
        previous = (firm, year - 1)
        adjacent = complete and previous in panel and panel[previous]['complete'] == '1'
        row = [firm, year, len(labels) if complete else None,
               len(labels - hits[previous]) if adjacent else None,
               len(hits[previous] - labels) if adjacent else None,
               sum(earliest[p] == year for p in labels) if complete else None]
        if adjacent and len(labels) - len(hits[previous]) != row[3] - row[4]:
            raise ValueError('Entry/exit identity failed')
        records.append({'firm': firm, 'year': year, 'complete': complete,
                        'labels': sorted(labels), 'panel': row})
    output.mkdir(parents=True)
    payload = {'industry': v['industry'], 'records': records,
               'normalized_labels': sorted({normalize(p) for labels in hits.values() for p in labels}),
               'earliest_years': earliest, 'world_metadata': metadata, 'embedding_contract': CONTRACT,
               'source_hashes': {str(p.resolve()): sha(p) for p in
                   [taxonomy, original, assignments, final/'validation.json', coverage, world, world_metadata]}}
    dump(output / 'payload.json', payload)
    return payload

def embed(directory, maximum):
    """Bounded paid requests; per-batch intents prevent blind replay after an uncertain response."""
    payload = json.loads((directory / 'payload.json').read_text())
    cache = directory / 'embedding_cache'; cache.mkdir(exist_ok=True)
    labels = payload['normalized_labels']; vectors = {}; missing = []
    def path(label):
        key = hashlib.sha256(json.dumps([CONTRACT, label], sort_keys=True).encode()).hexdigest()
        return cache / (key + '.json')
    for label in labels:
        p = path(label)
        if p.exists():
            item = json.loads(p.read_text())
            if item['label'] != label or item['contract'] != CONTRACT:
                raise ValueError('Embedding cache contract mismatch')
            vectors[label] = item['vector']
        else:
            missing.append(label)
    if len(missing) > maximum:
        raise ValueError(f'{len(missing)} new labels exceed --max-new-labels={maximum}; no calls made')
    if missing and not os.environ.get('GEMINI_API_KEY'):
        raise ValueError('GEMINI_API_KEY is required')
    for start in range(0, len(missing), 32):
        chunk = missing[start:start+32]
        batch_id = hashlib.sha256(json.dumps([CONTRACT, chunk], sort_keys=True).encode()).hexdigest()
        intent = cache / (batch_id + '.intent')
        # Reconciliation is per label, so changing batch boundaries cannot replay uncertain work.
        for old in cache.glob('*.intent'):
            if set(json.loads(old.read_text())['labels']) & set(chunk):
                raise RuntimeError('Uncertain embedding request; reconcile its .intent before retrying')
        dump(intent, {'labels': chunk, 'contract': CONTRACT, 'status': 'submission_intent'})
        body = {'requests': [{'model': 'models/' + CONTRACT['model'],
                 'content': {'parts': [{'text': CONTRACT['prefix'] + label}]},
                 'outputDimensionality': CONTRACT['dimensions']} for label in chunk]}
        req = urllib.request.Request('https://generativelanguage.googleapis.com/v1beta/models/' +
            CONTRACT['model'] + ':batchEmbedContents', data=json.dumps(body).encode(),
            headers={'Content-Type': 'application/json', 'x-goog-api-key': os.environ['GEMINI_API_KEY']})
        with urllib.request.urlopen(req, timeout=120) as response:
            result = json.load(response)
        values = [item['values'] for item in result['embeddings']]
        array = np.asarray(values, dtype=float)
        if array.shape != (len(chunk), 3072) or not np.isfinite(array).all() or (np.linalg.norm(array,axis=1)<1e-12).any():
            raise ValueError('Invalid embedding response; preserve request intent for review')
        for label, vector in zip(chunk, values):
            dump(path(label), {'label': label, 'vector': vector, 'contract': CONTRACT})
            vectors[label] = vector
        intent.unlink()
    array = np.asarray([vectors[label] for label in labels], dtype=float).reshape(len(labels), 3072)
    np.save(directory / 'vectors.npy', array)
    dump(directory / 'vectors.json', {'contract': CONTRACT, 'labels': labels,
        'payload_sha256': sha(directory/'payload.json'), 'vectors_sha256': sha(directory/'vectors.npy')})

def calculate(payload, raw):
    records, labels = payload['records'], payload['normalized_labels']
    if raw.shape != (len(labels), 3072) or not np.isfinite(raw).all():
        raise ValueError('Invalid embedding dimensions or values')
    norms = np.linalg.norm(raw, axis=1)
    if (norms < 1e-12).any():
        raise ValueError('Zero raw embedding')
    unit = raw / norms[:, None] if len(labels) else raw
    centered = unit - unit.mean(axis=0) if len(labels) else unit
    index = {label:i for i,label in enumerate(labels)}
    label_sets = [set(map(normalize, r['labels'])) for r in records]
    n = sum(r['complete'] for r in records)
    df = collections.Counter(label for r,ls in zip(records,label_sets) if r['complete'] for label in ls)
    weights = np.array([np.sqrt(1 + np.log((n+1)/(df[label]+1))) for label in labels])
    vectors = []
    for r,ls in zip(records,label_sets):
        ids = [index[label] for label in sorted(ls)]
        v = (weights[ids,None] * centered[ids]).sum(axis=0)
        norm = np.linalg.norm(v)
        vectors.append(v/norm if r['complete'] and ids and norm >= 1e-12 else None)
    panel = []
    for i,r in enumerate(records):
        peers = []
        if vectors[i] is not None:
            for j,other in enumerate(records):
                if other['firm'] == r['firm'] or other['year'] != r['year'] or vectors[j] is None:
                    continue
                distance = 1-float(np.clip(vectors[i] @ vectors[j], -1, 1))
                check = .5*float(np.sum((vectors[i]-vectors[j])**2))
                if abs(distance-check) > 1e-10:
                    raise ValueError('Cosine/Euclidean identity failed')
                peers.append((distance, other['firm']))
        peers.sort()
        panel.append(r['panel'] + ([peers[0][1], peers[0][0], float(np.mean([x[0] for x in peers]))]
                                  if peers else [None,None,None]))
    return panel

def build(directory):
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.comments import Comment
    payload = json.loads((directory/'payload.json').read_text())
    meta = json.loads((directory/'vectors.json').read_text())
    if (meta['contract'] != CONTRACT or meta['labels'] != payload['normalized_labels'] or
        meta['payload_sha256'] != sha(directory/'payload.json') or
        meta['vectors_sha256'] != sha(directory/'vectors.npy')):
        raise ValueError('Embedding provenance mismatch')
    panel = calculate(payload,np.load(directory/'vectors.npy',allow_pickle=False))
    workbook = Workbook(); workbook.remove(workbook.active)
    records = payload['records']; titles = set(); expected_sheets = {}
    for firm in sorted({r['firm'] for r in records}):
        title = re.sub(r'[\\/*?:\[\]]','_',firm)[:31]
        if not title or title.casefold() in titles or title.casefold() == 'firm year panel':
            raise ValueError('Firm IDs do not map uniquely to valid Excel sheet names')
        titles.add(title.casefold())
        if title != firm:
            raise ValueError('Use canonical firm IDs that are valid Excel sheet names (<=31 characters)')
        rows = [r for r in records if r['firm']==firm]
        labels = sorted({p for r in rows for p in r['labels']},key=lambda x:(x.casefold(),x))
        values = [['Parent label']+[r['year'] for r in rows]]
        values.extend([[label]+[int(label in r['labels']) if r['complete'] else 0 for r in rows] for label in labels])
        ws=workbook.create_sheet(title);expected_sheets[title]=values
        for row in values:ws.append(row)
        ws.freeze_panes='B2';ws.column_dimensions['A'].width=96
        for j,r in enumerate(rows,2):
            ws.column_dimensions[ws.cell(1,j).column_letter].width=10
            if not r['complete']:
                for i in range(1,len(labels)+2):
                    ws.cell(i,j).comment=Comment('Missing assignments: zero placeholders are not observed absences.','SystemAct')
        for row in ws.iter_rows(min_row=2,min_col=2):
            for cell in row:
                if cell.value==1:cell.fill=PatternFill('solid',fgColor='C6EFCE')
    ws=workbook.create_sheet('firm year panel');ws.append(COLUMNS)
    for row in panel:ws.append(row)
    expected_sheets[ws.title]=[COLUMNS]+panel;ws.freeze_panes='C2'
    for col,width in zip('ABCDEFGHI',[25,10,18,18,18,24,25,25,25]):ws.column_dimensions[col].width=width
    ws['F1'].comment=Comment('First observed in the declared available corpus, not invention. '+payload['world_metadata']['scope_description'],'SystemAct')
    for row in ws.iter_rows(min_row=2,min_col=8,max_col=9):
        for cell in row:cell.number_format='0.000000'
    for sheet in workbook:
        sheet.sheet_view.showGridLines=False
        for row in sheet:
            for cell in row:
                cell.font=Font(name='Arial',size=10,color='000000',bold=cell.row==1)
                if isinstance(cell.value,str):cell.data_type='s'
    path=directory/('matrix_'+payload['industry']+'.xlsx')
    if path.exists():raise FileExistsError('Matrix already exists; choose a fresh build directory')
    workbook.save(path)
    saved=load_workbook(path,data_only=True)
    for name,expected in expected_sheets.items():
        actual=list(saved[name].values)
        for got,want in zip(actual,expected):
            for g,w in zip(got,want):
                if isinstance(w,float):
                    if abs(g-w)>1e-12:raise ValueError('Saved numeric value changed')
                elif g!=w:raise ValueError('Saved cell differs from expected value')
        if len(actual)!=len(expected):raise ValueError('Saved row count mismatch')
    with (directory/'firm_year_panel.csv').open('w',newline='') as f:
        writer=csv.writer(f);writer.writerow(COLUMNS);writer.writerows(panel)
    dump(directory/'validation.json',{'pass':True,'industry':payload['industry'],
         'firm_years':len(panel),'firms':len(saved.sheetnames)-1,'workbook_sha256':sha(path),
         'numeric_and_cell_validation':True,'visual_render_validation':False,
         'source_hashes':payload['source_hashes'],'world_scope':payload['world_metadata']})
    return path

def main():
    parser=argparse.ArgumentParser(description=__doc__);sub=parser.add_subparsers(dest='stage',required=True)
    p=sub.add_parser('prepare')
    for name in ['run','coverage','world','world-metadata','output']:p.add_argument('--'+name,type=Path,required=True)
    p=sub.add_parser('embed');p.add_argument('--directory',type=Path,required=True)
    p.add_argument('--max-new-labels',type=int,required=True)
    p=sub.add_parser('build');p.add_argument('--directory',type=Path,required=True)
    a=parser.parse_args()
    if a.stage=='prepare':prepare(a.run,a.coverage,a.world,a.world_metadata,a.output)
    elif a.stage=='embed':embed(a.directory,a.max_new_labels)
    else:print(build(a.directory))

if __name__=='__main__':main()
